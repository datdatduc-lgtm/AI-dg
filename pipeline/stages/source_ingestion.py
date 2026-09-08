"""Source Ingestion (M2).

Builds a canonical Source Package from a 2D source file (PDF / DXF / Excel /
SketchUp linework) as described in `prompt/04_2D_SOURCE_INGESTION.md`.

Every extracted fact keeps provenance: source_id, page/sheet/layer, original
text and confidence.  No closed polyline is interpreted as a panel here; this
module only extracts candidates for later reconciliation.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

PROJECT_ROOT = Path(__file__).resolve().parents[2]
OPTIONAL_DEPENDENCY_DIR = PROJECT_ROOT / "OUTPUT" / "pipeline_deps"
if OPTIONAL_DEPENDENCY_DIR.is_dir() and str(OPTIONAL_DEPENDENCY_DIR) not in sys.path:
    sys.path.insert(0, str(OPTIONAL_DEPENDENCY_DIR))

try:
    import fitz  # PyMuPDF
except Exception:  # pragma: no cover - optional dependency
    fitz = None

try:
    import ezdxf
except Exception:  # pragma: no cover - optional dependency
    ezdxf = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional dependency
    PdfReader = None

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - optional dependency
    load_workbook = None

try:
    from PIL import Image
except Exception:  # pragma: no cover - optional dependency
    Image = None

try:
    import pytesseract
except Exception:  # pragma: no cover - optional dependency
    pytesseract = None

DRAWING_TYPES = [
    "PLAN",
    "ELEVATION",
    "SECTION",
    "DETAIL",
    "SCHEDULE",
    "MATERIAL_LEGEND",
    "UNKNOWN",
]

DIMENSION_RE = re.compile(
    r"(?<!\d)(\d{2,5}(?:\.\d+)?)\s*[xX×]\s*(\d{2,5}(?:\.\d+)?)"
    r"(?:\s*[xX×]\s*(\d{1,4}(?:\.\d+)?))?"
)
MATERIAL_RE = re.compile(
    r"\b(?:MDF|HDF|MFC|PLY|PLYWOOD|WD|LAM|VEN|VENEER|HW|AL|KÍNH|KINH|NHỰA|NHUA|INOX"
    r"|ĐÁ|DA|MIKA)[-_ ]?\d{0,4}\b",
    re.IGNORECASE,
)
NOTE_RE = re.compile(r"(?:GHP|GHI CHÚ|NOTE|CHÚ THÍCH|CHU THICH)\s*[:;.-]?\s*(.+)", re.IGNORECASE)
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


@dataclass
class SourcePackage:
    """Canonical extracted package for one input source."""

    source_id: str
    source_type: str
    path: str
    pages: list[dict[str, Any]] = field(default_factory=list)
    drawings: list[dict[str, Any]] = field(default_factory=list)
    dimensions: list[dict[str, Any]] = field(default_factory=list)
    notes: list[dict[str, Any]] = field(default_factory=list)
    materials: list[dict[str, Any]] = field(default_factory=list)
    confidence: dict[str, float] = field(default_factory=dict)
    schema_version: str = "0.1"
    generated_utc: str = field(default_factory=utcnow)

    def to_dict(self) -> dict[str, Any]:
        return {
            "schema_version": self.schema_version,
            "source_id": self.source_id,
            "type": self.source_type,
            "path": self.path,
            "pages": self.pages,
            "drawings": self.drawings,
            "dimensions": self.dimensions,
            "notes": self.notes,
            "materials": self.materials,
            "confidence": self.confidence,
            "generated_utc": self.generated_utc,
        }


def new_source_id(index: int) -> str:
    return f"SRC-{index:03d}"


def classify_drawing(text: str) -> dict[str, Any]:
    """Classify page text into a drawing type with confidence."""
    upper = text.upper()
    keywords: dict[str, list[str]] = {
        "PLAN": ["MẶT BẰNG", "MAT BANG", "PLAN", "TOP VIEW", "BỐ TRÍ", "BO TRI"],
        "ELEVATION": ["MẶT ĐỨNG", "MAT DUNG", "MẶT BÊN", "MAT BEN", "ELEVATION", "FRONT"],
        "SECTION": ["MẶT CẮT", "MAT CAT", "SECTION", "CUT A-A", "CẮT"],
        "DETAIL": ["CHI TIẾT", "CHI TIET", "DETAIL", "CT1", "CT2"],
        "SCHEDULE": ["THỐNG KÊ", "THONG KE", "SCHEDULE", "BẢNG", "DANH MỤC", "LIST"],
        "MATERIAL_LEGEND": ["GHI CHÚ", "GHI CHU", "LEGEND", "CHÚ THÍCH", "VẬT LIỆU", "VAT LIEU"],
    }
    hits = []
    for d_type, kws in keywords.items():
        count = sum(1 for kw in kws if kw in upper)
        if count:
            hits.append((count, d_type))
    if not hits:
        return {"drawing_type": "UNKNOWN", "confidence": 0.0, "signals": []}
    best = max(hits, key=lambda x: x[0])
    return {
        "drawing_type": best[1],
        "confidence": round(min(1.0, 0.45 + 0.15 * best[0]), 3),
        "signals": [s for _, s in sorted(hits, reverse=True)],
    }


def _dimension_values(raw: str) -> list[float]:
    """Extract a bounded list of numeric values from one dimension label."""
    values: list[float] = []
    for part in re.split(r"[xX×]", raw):
        match = re.search(r"\d{1,6}(?:\.\d+)?", part)
        if match:
            values.append(float(match.group(0)))
    return values[:3]


def _text_facts(text: str, page: int = 1) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """Extract dimension/material/note candidates from OCR or vector text."""
    dimensions: list[dict[str, Any]] = []
    materials: list[dict[str, Any]] = []
    notes: list[dict[str, Any]] = []
    for match in DIMENSION_RE.finditer(text):
        values = _dimension_values(match.group(0))
        dimensions.append({
            "raw": match.group(0),
            "values_mm": values,
            "a_mm": values[0] if values else None,
            "b_mm": values[1] if len(values) > 1 else None,
            "c_mm": values[2] if len(values) > 2 else None,
            "page": page,
            "provenance": {"page": page, "original_text": match.group(0)},
            "confidence": 0.9,
        })
    for code in sorted({m.group(0).strip() for m in MATERIAL_RE.finditer(text)}):
        materials.append({
            "code": code,
            "pages": [page],
            "provenance": {"page": page, "original_text": code},
            "confidence": 0.8,
        })
    for match in NOTE_RE.finditer(text):
        notes.append({
            "text": match.group(1).strip(),
            "page": page,
            "provenance": {"page": page, "original_text": match.group(0)},
            "confidence": 0.7,
        })
    return dimensions, materials, notes


def _resolve_tesseract_executable() -> str | None:
    """Find an explicitly installed Tesseract binary without downloading it."""
    if pytesseract is None:
        return None
    configured = str(getattr(getattr(pytesseract, "pytesseract", None), "tesseract_cmd", "") or "")
    candidates = []
    if configured and configured.lower() != "tesseract":
        candidates.append(Path(configured))
    candidates.extend([
        PROJECT_ROOT / "OUTPUT" / "tesseract" / "tesseract.exe",
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    ])
    for candidate in candidates:
        if candidate.is_file():
            return str(candidate)
    return shutil.which("tesseract")


def _ocr_image(image: Any) -> tuple[str, str, float]:
    """Run local OCR only when a binary is already present; never fetch one."""
    if pytesseract is None:
        return "", "NOT_AVAILABLE", 0.0
    executable = _resolve_tesseract_executable()
    if not executable:
        return "", "BINARY_NOT_FOUND", 0.0
    original_command = getattr(pytesseract.pytesseract, "tesseract_cmd", "tesseract")
    try:
        pytesseract.pytesseract.tesseract_cmd = executable
        text = pytesseract.image_to_string(image, config="--psm 11")[:100_000]
        confidence = 0.0
        try:
            data = pytesseract.image_to_data(image, config="--psm 11", output_type=pytesseract.Output.DICT)
            values = [float(value) for value in data.get("conf", []) if str(value).strip() not in {"", "-1"} and float(value) >= 0]
            confidence = round(sum(values) / len(values) / 100.0, 3) if values else (0.5 if text.strip() else 0.0)
        except Exception:
            confidence = 0.5 if text.strip() else 0.0
        return text, "OK" if text.strip() else "EMPTY", confidence
    except Exception as exc:
        return "", f"ERROR_{exc.__class__.__name__}", 0.0
    finally:
        pytesseract.pytesseract.tesseract_cmd = original_command


def _annotate_ocr_facts(
    dimensions: list[dict[str, Any]],
    materials: list[dict[str, Any]],
    page: int,
    image_path: Path,
    confidence: float,
) -> None:
    """Keep OCR candidates reviewable instead of silently treating them as exact facts."""
    candidate_confidence = round(min(0.7, max(0.25, confidence * 0.8)), 3)
    for fact in dimensions:
        fact["confidence"] = candidate_confidence
        fact["state"] = "OCR_CANDIDATE"
        fact["provenance"] = {
            **(fact.get("provenance") or {}),
            "page": page,
            "image_path": str(image_path),
            "extraction": "ocr",
        }
    for fact in materials:
        fact["confidence"] = candidate_confidence
        fact["state"] = "OCR_CANDIDATE"
        fact["provenance"] = {
            **(fact.get("provenance") or {}),
            "page": page,
            "image_path": str(image_path),
            "extraction": "ocr",
        }


def _ingest_pdf_with_pypdf(path: Path) -> SourcePackage:
    """Portable PDF text fallback when PyMuPDF is unavailable."""
    if PdfReader is None:
        raise RuntimeError("PyMuPDF or pypdf is required for PDF ingestion")
    reader = PdfReader(str(path))
    pages: list[dict[str, Any]] = []
    dimensions: list[dict[str, Any]] = []
    materials: list[dict[str, Any]] = []
    notes: list[dict[str, Any]] = []
    for index, page in enumerate(reader.pages):
        page_number = index + 1
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        classification = classify_drawing(text)
        page_dimensions, page_materials, page_notes = _text_facts(text, page=page_number)
        dimensions.extend(page_dimensions)
        materials.extend(page_materials)
        notes.extend(page_notes)
        pages.append({
            "page": page_number,
            "text_chars": len(text),
            "text_excerpt": text[:100_000],
            "likely_scanned_or_image_only": len(text) < 20,
            "classification": classification,
            "title_candidates": _title_candidates(text),
            "rendered_image": None,
            "render_status": "UNAVAILABLE_WITH_PYPDF",
        })
    return SourcePackage(
        source_id="",
        source_type="pdf",
        path=str(path),
        pages=pages,
        drawings=[{"page": p["page"], **p["classification"]} for p in pages],
        dimensions=dimensions,
        notes=notes,
        materials=materials,
        confidence={
            "text_extraction": 0.8 if any(p["text_chars"] for p in pages) else 0.2,
            "dimension_extraction": 0.9,
            "material_extraction": 0.8,
            "drawing_classification": round(sum(p["classification"]["confidence"] for p in pages) / max(1, len(pages)), 3),
            "rendering": 0.0,
        },
    )
def _title_candidates(text: str) -> list[str]:
    """Heuristic title block lines (VACH NGAN / TU / QUAY ...)."""
    candidates = []
    for line in text.splitlines():
        stripped = line.strip()
        if len(stripped) < 4:
            continue
        for kw in ("VÁCH NGĂN", "VACH NGAN", "TỦ", "TU ", "KỆ", "KE ", "QUẦY", "QUAY", "BÀN", "BAN "):
            if kw in stripped.upper():
                candidates.append(stripped[:80])
                break
    return candidates[:10]


def ingest_pdf(path: Path, render: bool = False, work_dir: Optional[Path] = None) -> SourcePackage:
    """Extract a Source Package from a PDF.

    Vector text is preferred; scanned/image-only pages are flagged so they can
    be reviewed visually instead of being silently trusted.
    """
    path = path.expanduser().resolve()
    if not path.is_file() or path.suffix.lower() != ".pdf":
        raise ValueError(f"PDF not found or invalid extension: {path}")
    if fitz is None:
        return _ingest_pdf_with_pypdf(path)

    doc = fitz.open(path)
    pages: list[dict[str, Any]] = []
    dimensions: list[dict[str, Any]] = []
    notes: list[dict[str, Any]] = []
    materials_map: dict[str, dict[str, Any]] = {}
    scanned_pages: list[int] = []

    for index, page in enumerate(doc):
        page_number = index + 1
        text = page.get_text("text").strip()
        likely_scanned = len(text) < 20
        if likely_scanned:
            scanned_pages.append(page_number)

        classification = classify_drawing(text)

        page_dimensions, page_materials_facts, page_notes = _text_facts(text, page=page_number)
        dimensions.extend(page_dimensions)
        notes.extend(page_notes)

        for material in page_materials_facts:
            code = material["code"]
            entry = materials_map.setdefault(code, {"code": code, "confidence": 0.8})
            entry.setdefault("pages", set()).add(page_number)
            entry.setdefault("provenance", {"page": page_number, "original_text": code})

        image_path = None
        ocr_text = ""
        ocr_status = "NOT_REQUESTED"
        ocr_confidence = 0.0
        if render:
            zoom = 2.0  # 144 dpi
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            pages_dir = (work_dir or path.parent) / "extracted" / "pages"
            pages_dir.mkdir(parents=True, exist_ok=True)
            image_path = pages_dir / f"page-{page_number:04d}.png"
            pix.save(image_path)
            if likely_scanned and Image is not None:
                try:
                    with Image.open(image_path) as rendered_image:
                        ocr_text, ocr_status, ocr_confidence = _ocr_image(rendered_image)
                    if ocr_text.strip():
                        ocr_dimensions, ocr_materials, ocr_notes = _text_facts(ocr_text, page=page_number)
                        _annotate_ocr_facts(ocr_dimensions, ocr_materials, page_number, image_path, ocr_confidence)
                        dimensions.extend(ocr_dimensions)
                        notes.extend(ocr_notes)
                        for material in ocr_materials:
                            code = material["code"]
                            entry = materials_map.setdefault(code, {"code": code, "confidence": material.get("confidence", 0.5)})
                            entry["confidence"] = min(float(entry.get("confidence", 0.5)), float(material.get("confidence", 0.5)))
                            entry.setdefault("pages", set()).add(page_number)
                            entry.setdefault("provenance", material.get("provenance"))
                        if not text:
                            classification = classify_drawing(ocr_text)
                except Exception as exc:
                    ocr_status = f"ERROR_{exc.__class__.__name__}"

        pages.append(
            {
                "page": page_number,
                "text_chars": len(text),
                "text_excerpt": text[:100_000],
                "likely_scanned_or_image_only": likely_scanned,
                "classification": classification,
                "title_candidates": _title_candidates(text or ocr_text),
                "rendered_image": str(image_path) if image_path else None,
                "ocr_status": ocr_status,
                "ocr_text_chars": len(ocr_text),
                "ocr_text_excerpt": ocr_text[:100_000],
                "ocr_confidence": ocr_confidence,
                "text_source": "pdf_text" if text else ("ocr" if ocr_text else "none"),
            }
        )

    doc.close()

    confidence = {
        "text_extraction": 0.3 if scanned_pages else 0.9,
        "ocr": round(sum(float(p.get("ocr_confidence", 0.0)) for p in pages) / max(1, len(pages)), 3),
        "dimension_extraction": 0.9,
        "material_extraction": 0.8,
        "drawing_classification": round(
            sum(p["classification"]["confidence"] for p in pages) / max(1, len(pages)), 3
        ),
    }

    materials: list[dict[str, Any]] = []
    for code, entry in materials_map.items():
        entry["pages"] = sorted(int(p) for p in entry.pop("pages", set()))
        materials.append(entry)

    return SourcePackage(
        source_id="",
        source_type="pdf",
        path=str(path),
        pages=pages,
        drawings=[{"page": p["page"], **p["classification"]} for p in pages],
        dimensions=dimensions,
        notes=notes,
        materials=materials,
        confidence=confidence,
    )


def ingest_dxf(path: Path) -> SourcePackage:
    """Extract a Source Package from DXF (read-only via ezdxf).

    Reads layers, polylines (closed profiles), texts, dimensions and blocks.
    Does not auto-interpret every closed polyline as a panel.
    """
    path = path.expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"File not found: {path}")

    if ezdxf is None:
        return _ingest_dxf_ascii(path)

    try:
        doc = ezdxf.readfile(path)
    except Exception:
        # Some exporters omit optional subclass markers. Preserve the source
        # as a read-only candidate through the bounded ASCII fallback instead
        # of dropping the complete package.
        return _ingest_dxf_ascii(path)
    msp = doc.modelspace()
    layers = [layer.dxf.name for layer in doc.layers]
    entity_summary: dict[str, int] = {}
    texts: list[str] = []
    closed_profiles = 0
    dims: list[dict[str, Any]] = []
    blocks: list[str] = []

    for entity in msp:
        etype = entity.dxftype()
        entity_summary[etype] = entity_summary.get(etype, 0) + 1
        if etype in ("TEXT", "MTEXT"):
            texts.append(str(entity.dxf.text))
        elif etype == "LWPOLYLINE" and entity.closed:
            closed_profiles += 1
            points = [(float(point[0]), float(point[1])) for point in entity.get_points()]
            dims.extend(_profile_span_dimensions(points, entity.dxf.layer, "LWPOLYLINE"))
        elif etype == "POLYLINE" and entity.is_closed:
            closed_profiles += 1
        elif etype in ("DIMENSION", "DIMENSION_ALIGNED", "DIMENSION_LINEAR"):
            raw_value = getattr(entity.dxf, "actual_measurement", None)
            dims.append({
                "type": etype,
                "layer": entity.dxf.layer,
                "value_mm": float(raw_value) if isinstance(raw_value, (int, float)) else None,
                "provenance": {"layer": entity.dxf.layer, "entity_type": etype},
                "confidence": 0.85 if isinstance(raw_value, (int, float)) else 0.4,
            })
        elif etype == "INSERT":
            blocks.append(str(getattr(entity.dxf, "name", "")))

    text_payload = "\n".join(texts)
    classification = classify_drawing(text_payload)
    material_codes = sorted({m.group(0).strip() for m in MATERIAL_RE.finditer(text_payload)})

    return SourcePackage(
        source_id="",
        source_type="dxf",
        path=str(path),
        pages=[
            {
                "page": 1,
                "text_chars": sum(len(t) for t in texts),
                "likely_scanned_or_image_only": False,
                "classification": classification,
                "layers": layers,
                "entity_summary": entity_summary,
                "closed_profiles": closed_profiles,
                "blocks": sorted(set(blocks)),
                "title_candidates": _title_candidates(text_payload),
            }
        ],
        dimensions=dims,
        materials=[{"code": c, "provenance": {"layer": "any"}, "confidence": 0.7} for c in material_codes],
        confidence={
            "cad_parse": 0.9,
            "dimension_extraction": 0.85,
            "material_extraction": 0.7,
            "drawing_classification": classification["confidence"],
        },
    )
def _normalized_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _numeric_mm(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"(?<!\d)(\d+(?:\.\d+)?)", str(value).replace(",", "."))
    return float(match.group(1)) if match else None


def ingest_excel(path: Path, sheet_name: Optional[str] = None) -> SourcePackage:
    """Extract a Source Package from an Excel schedule/table."""
    if load_workbook is None:
        raise RuntimeError("openpyxl is required for Excel ingestion")
    path = path.expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"File not found: {path}")

    workbook = load_workbook(path, data_only=True)
    sheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
    headers: list[str] = []
    rows: list[dict[str, Any]] = []
    dimensions: list[dict[str, Any]] = []
    materials: list[dict[str, Any]] = []
    notes: list[dict[str, Any]] = []
    for row in sheet.iter_rows(values_only=True):
        if not any(cell is not None for cell in row):
            continue
        if not headers:
            headers = [str(cell or f"COL_{i}").strip() for i, cell in enumerate(row)]
            continue
        rows.append({headers[i] if i < len(headers) else f"COL_{i}": cell for i, cell in enumerate(row)})

    normalized = {_normalized_header(header): header for header in headers}
    item_header = next((header for key, header in normalized.items() if any(token in key for token in ("item_code", "item", "code", "sku", "ma_hang", "ma"))), None)
    part_header = next((header for key, header in normalized.items() if any(token in key for token in ("part_id", "part", "component", "chi_tiet", "bo_phan"))), None)
    role_header = next((header for key, header in normalized.items() if any(token in key for token in ("part_role", "role", "vai_tro", "bo_phan"))), None)
    dim_headers = {
        "width": next((header for key, header in normalized.items() if key in {"width", "width_mm", "w", "rong", "rong_mm"} or "width" in key or "rong" in key), None),
        "depth": next((header for key, header in normalized.items() if key in {"depth", "depth_mm", "d", "sau", "sau_mm"} or "depth" in key or "sau" in key), None),
        "height": next((header for key, header in normalized.items() if key in {"height", "height_mm", "h", "cao", "cao_mm"} or "height" in key or "cao" in key), None),
    }
    material_headers = [header for key, header in normalized.items() if any(token in key for token in ("material", "vat_lieu", "mat"))]
    note_headers = [header for key, header in normalized.items() if any(token in key for token in ("note", "ghi_chu", "remark", "description", "mo_ta"))]
    for row_number, row in enumerate(rows, start=2):
        item_code = str(row.get(item_header) or "UNKNOWN") if item_header else "UNKNOWN"
        part_id = str(row.get(part_header) or "").strip() if part_header else ""
        part_role = str(row.get(role_header) or "").strip() if role_header else ""
        row_materials = [str(row.get(header)).strip() for header in material_headers if row.get(header) not in (None, "")]
        for dim, header in dim_headers.items():
            value_mm = _numeric_mm(row.get(header)) if header else None
            if value_mm is None:
                continue
            dimensions.append({
                "item_code": item_code,
                "part_id": part_id or None,
                "part_role": part_role or None,
                "material_code": row_materials[0] if row_materials else None,
                "dim": dim,
                "value_mm": value_mm,
                "axis": {"width": "X", "depth": "Y", "height": "Z"}[dim],
                "span": f"sheet-{sheet.title}-row-{row_number}",
                "row": row_number,
                "provenance": {"sheet": sheet.title, "row": row_number, "column": header, "original_value": row.get(header)},
                "confidence": 0.95,
            })
        for header in material_headers:
            value = row.get(header)
            if value not in (None, ""):
                materials.append({
                    "item_code": item_code,
                    "code": str(value).strip(),
                    "role": part_role or "unspecified",
                    "provenance": {"sheet": sheet.title, "row": row_number, "column": header, "original_value": value},
                    "confidence": 0.9,
                })
        for header in note_headers:
            value = row.get(header)
            if value not in (None, ""):
                notes.append({
                    "item_code": item_code,
                    "text": str(value).strip(),
                    "provenance": {"sheet": sheet.title, "row": row_number, "column": header},
                    "confidence": 0.8,
                })

    return SourcePackage(
        source_id="",
        source_type="excel",
        path=str(path),
        pages=[{"page": 1, "sheet": sheet.title, "rows": rows}],
        dimensions=dimensions,
        materials=materials,
        notes=notes,
        confidence={"excel_read": 0.95, "dimension_extraction": 0.95 if dimensions else 0.4},
    )


def ingest_image(path: Path, work_dir: Optional[Path] = None) -> SourcePackage:
    """Read an image source and optionally OCR it; never mutates the image."""
    if Image is None:
        raise RuntimeError("Pillow is required for image ingestion")
    path = path.expanduser().resolve()
    if not path.is_file() or path.suffix.lower() not in IMAGE_EXTENSIONS:
        raise ValueError(f"Image not found or invalid extension: {path}")
    with Image.open(path) as image:
        width, height = image.size
        ocr_text, ocr_status, ocr_confidence = _ocr_image(image)
        dimensions, materials, notes = _text_facts(ocr_text, page=1)
        _annotate_ocr_facts(dimensions, materials, page=1, image_path=path, confidence=ocr_confidence)
        classification = classify_drawing(ocr_text)
        page = {
            "page": 1,
            "width_px": width,
            "height_px": height,
            "text_chars": len(ocr_text),
            "ocr_status": ocr_status,
            "ocr_confidence": ocr_confidence,
            "ocr_text_excerpt": ocr_text[:100_000],
            "text_source": "ocr" if ocr_text else "none",
            "likely_scanned_or_image_only": True,
            "classification": classification,
            "title_candidates": _title_candidates(ocr_text),
        }
    return SourcePackage(
        source_id="",
        source_type="image",
        path=str(path),
        pages=[page],
        drawings=[{"page": 1, **classification}],
        dimensions=dimensions,
        materials=materials,
        notes=notes,
        confidence={
            "image_read": 0.95,
            "ocr": ocr_confidence if ocr_status == "OK" else 0.0,
            "dimension_extraction": 0.75 if dimensions else 0.0,
            "drawing_classification": classification["confidence"],
        },
    )


def _ingest_dxf_ascii(path: Path) -> SourcePackage:
    """Small read-only ASCII DXF fallback for layers/text/dimension candidates."""
    raw_lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    pairs: list[tuple[str, str]] = []
    for index in range(0, len(raw_lines) - 1, 2):
        pairs.append((raw_lines[index].strip(), raw_lines[index + 1].strip()))
    entities: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    in_entities = False
    for code, value in pairs:
        if code == "0":
            if current and current.get("type"):
                entities.append(current)
            current = {"type": value, "attrs": {}}
            if value == "SECTION":
                in_entities = False
            continue
        if current is None:
            continue
        if current.get("type") == "2" and value == "ENTITIES":
            in_entities = True
        attrs = current.setdefault("attrs", {})
        attrs.setdefault(code, []).append(value)
        if current.get("type") == "SECTION" and code == "2" and value == "ENTITIES":
            in_entities = True
    if current and current.get("type"):
        entities.append(current)

    entity_summary: dict[str, int] = {}
    layers: set[str] = set()
    texts: list[str] = []
    closed_profiles = 0
    blocks: set[str] = set()
    dims: list[dict[str, Any]] = []
    for entity in entities:
        etype = str(entity.get("type"))
        attrs = entity.get("attrs", {})
        if etype in {"SECTION", "ENDSEC", "EOF"}:
            continue
        entity_summary[etype] = entity_summary.get(etype, 0) + 1
        layer = (attrs.get("8") or [None])[0]
        if layer:
            layers.add(str(layer))
        if etype in {"TEXT", "MTEXT"}:
            texts.extend(str(value) for value in attrs.get("1", []))
        elif etype == "LWPOLYLINE" and str((attrs.get("70") or ["0"])[0]) in {"1", "129"}:
            closed_profiles += 1
            x_values = [float(value) for value in attrs.get("10", [])]
            y_values = [float(value) for value in attrs.get("20", [])]
            points = list(zip(x_values, y_values))
            dims.extend(_profile_span_dimensions(points, layer, "LWPOLYLINE"))
        elif etype == "POLYLINE" and str((attrs.get("70") or ["0"])[0]) in {"1", "129"}:
            closed_profiles += 1
        elif etype == "INSERT":
            blocks.update(str(value) for value in attrs.get("2", []))
        elif etype == "DIMENSION":
            raw_value = (attrs.get("42") or [None])[0]
            try:
                value_mm = float(raw_value) if raw_value is not None else None
            except ValueError:
                value_mm = None
            dims.append({
                "type": etype,
                "layer": layer,
                "value_mm": value_mm,
                "provenance": {"layer": layer, "entity_type": etype},
                "confidence": 0.75 if value_mm is not None else 0.35,
            })
    text_payload = "\n".join(texts)
    classification = classify_drawing(text_payload)
    material_codes = sorted({m.group(0).strip() for m in MATERIAL_RE.finditer(text_payload)})
    return SourcePackage(
        source_id="",
        source_type="dxf",
        path=str(path),
        pages=[{
            "page": 1,
            "text_chars": len(text_payload),
            "likely_scanned_or_image_only": False,
            "classification": classification,
            "layers": sorted(layers),
            "entity_summary": entity_summary,
            "closed_profiles": closed_profiles,
            "blocks": sorted(blocks),
            "title_candidates": _title_candidates(text_payload),
            "parser": "ascii_fallback",
        }],
        dimensions=dims,
        materials=[{"code": code, "provenance": {"layer": "any"}, "confidence": 0.7} for code in material_codes],
        confidence={
            "cad_parse": 0.65,
            "dimension_extraction": 0.75 if dims else 0.3,
            "material_extraction": 0.7,
            "drawing_classification": classification["confidence"],
        },
    )


def _profile_span_dimensions(
    points: list[tuple[float, float]], layer: str | None, entity_type: str
) -> list[dict[str, Any]]:
    """Return geometry candidates for a closed profile without semantic inference."""
    if len(points) < 2:
        return []
    x_values = [point[0] for point in points]
    y_values = [point[1] for point in points]
    spans = (("X", max(x_values) - min(x_values)), ("Y", max(y_values) - min(y_values)))
    return [
        {
            "fact": "closed profile span candidate",
            "axis": axis,
            "span": f"{layer or 'UNLAYERED'}:{entity_type}",
            "value_mm": round(float(value), 3),
            "provenance": {
                "layer": layer,
                "entity_type": entity_type,
                "evidence": "closed_profile_bounding_span",
                "point_count": len(points),
            },
            "confidence": 0.65,
        }
        for axis, value in spans
        if value > 0
    ]


def ingest_skp_metadata(path: Path) -> SourcePackage:
    """Record safe file metadata for an SKP; geometry stays in SketchUp API."""
    path = path.expanduser().resolve()
    if not path.is_file() or path.suffix.lower() != ".skp":
        raise ValueError(f"SKP not found or invalid extension: {path}")
    stat = path.stat()
    return SourcePackage(
        source_id="",
        source_type="skp",
        path=str(path),
        pages=[{
            "page": 1,
            "file_name": path.name,
            "size_bytes": stat.st_size,
            "adapter_status": "OFFICIAL_SKETCHUP_API_REQUIRED",
            "geometry_read": False,
            "likely_scanned_or_image_only": False,
            "classification": {"drawing_type": "UNKNOWN", "confidence": 0.0, "signals": []},
        }],
        confidence={"skp_metadata_read": 0.95, "geometry_read": 0.0},
    )


def ingest_source(path: Path, render: bool = False, work_dir: Optional[Path] = None) -> SourcePackage:
    """Dispatch one supported source without changing the source file."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return ingest_pdf(path, render=render, work_dir=work_dir)
    if suffix in {".dxf", ".dwg"}:
        return ingest_dxf(path)
    if suffix in {".xlsx", ".xls"}:
        return ingest_excel(path)
    if suffix in IMAGE_EXTENSIONS:
        return ingest_image(path, work_dir=work_dir)
    if suffix == ".skp":
        return ingest_skp_metadata(path)
    raise ValueError(f"Unsupported source type: {path.suffix}")


def save_package(package: SourcePackage, output_path: Path) -> Path:
    output_path = output_path.expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(package.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
    return output_path


def load_package(path: Path) -> SourcePackage:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    known = set(SourcePackage.__dataclass_fields__)
    return SourcePackage(**{k: v for k, v in data.items() if k in known})
